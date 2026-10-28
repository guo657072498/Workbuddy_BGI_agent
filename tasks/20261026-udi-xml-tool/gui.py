"""UDI XML 转 Excel —— 图形界面。

设计要点：
1. 转换在后台线程跑，界面不卡死，进度条实时更新。
2. 优先使用 Windows 原生拖拽（tkinterdnd2）；不可用时自动退回「添加文件/文件夹」按钮。
3. 所有异常都在界面里以弹窗或日志呈现，绝不静默退出。
"""

import os
import queue
import subprocess
import sys
import threading
import traceback

import tkinter as tk
from tkinter import filedialog, messagebox, ttk

# 打包后 _MEIPASS 是临时解包目录，源码运行时是脚本目录
if getattr(sys, "frozen", False):
    APP_DIR = os.path.dirname(sys.executable)
    RES_DIR = getattr(sys, "_MEIPASS", APP_DIR)
else:
    APP_DIR = os.path.dirname(os.path.abspath(__file__))
    RES_DIR = APP_DIR
sys.path.insert(0, RES_DIR)

from converter import convert          # noqa: E402
from udi_parser import source as src_mod  # noqa: E402
import masterdata                      # noqa: E402
from masterdata import store as md_store  # noqa: E402

# ---------------------------------------------------------------- 拖拽支持

DND_OK = False
try:
    from tkinterdnd2 import DND_FILES, TkinterDnD
    DND_OK = True
except Exception:
    DND_FILES = None


# ---------------------------------------------------------------- 主窗口

class App:
    def __init__(self, root):
        self.root = root
        self.sources = []          # XmlSource 列表
        self.queue = queue.Queue()
        self.worker = None
        self.cancel_flag = threading.Event()
        self.results = []
        self._md_win = {}
        self._build_ui()
        self._init_masterdata()
        self._poll_queue()

    # ------------------------------------------------------------ 菜单

    def _build_menu(self):
        menubar = tk.Menu(self.root)

        m_file = tk.Menu(menubar, tearoff=False)
        m_file.add_command(label="添加文件…", command=self._add_files)
        m_file.add_command(label="添加文件夹…", command=self._add_folder)
        m_file.add_separator()
        m_file.add_command(label="打开输出目录", command=self._open_out)
        m_file.add_separator()
        m_file.add_command(label="退出", command=self.root.destroy)
        menubar.add_cascade(label="文件", menu=m_file)

        m_md = tk.Menu(menubar, tearoff=False)
        m_md.add_command(label="产品主数据…（DI码 → 成品名称/规格）",
                         command=lambda: self._open_masterdata(md_store.KIND_PRODUCT))
        m_md.add_command(label="法人主数据…（法人编码 → 法人名称）",
                         command=lambda: self._open_masterdata(md_store.KIND_CORP))
        m_md.add_separator()
        m_md.add_command(label="打开主数据文件夹", command=self._open_md_folder)
        menubar.add_cascade(label="主数据", menu=m_md)

        m_help = tk.Menu(menubar, tearoff=False)
        m_help.add_command(label="使用说明", command=self._help)
        m_help.add_command(label="关于", command=self._about)
        menubar.add_cascade(label="帮助", menu=m_help)

        self.root.configure(menu=menubar)

    # ------------------------------------------------------------ 主数据

    def _init_masterdata(self):
        """启动时确保主数据就绪，并检查内置数据是否有新增。"""
        try:
            state = md_store.ensure_seeded()
        except Exception as e:
            self._log(f"主数据初始化失败：{e}", "error")
            return

        created = [k for k, v in state.items() if v == "created"]
        no_seed = [k for k, v in state.items() if v == "no_seed"]
        if created:
            self._log(f"已初始化主数据（{'、'.join(md_store.TITLES[k] for k in created)}），"
                      f"位置：{md_store.data_dir()}", "info")
        if no_seed:
            self._log(f"缺少内置主数据种子：{'、'.join(no_seed)}，请检查打包配置", "error")

        # 磁盘已有数据时，看看新版内置数据里有没有新条目，有就问要不要合并
        extra = {}
        for kind in (md_store.KIND_PRODUCT, md_store.KIND_CORP):
            if state.get(kind) != "kept":
                continue
            seed_only = md_store.seed_diff(kind)
            if not seed_only:
                continue
            pk = md_store.column_defs(kind)[0][0]
            try:
                current = {r[pk] for r in md_store.load_rows(kind)}
            except Exception:
                continue
            new_keys = [k for k in seed_only if k not in current]
            if new_keys:
                extra[kind] = (seed_only, new_keys)

        if extra:
            lines = [f"{md_store.TITLES[k]}：新增 {len(v[1])} 条" for k, v in extra.items()]
            if messagebox.askyesno(
                "主数据有新增",
                "检测到本版本内置的主数据比当前使用的主数据多出一些条目：\n\n"
                + "\n".join(lines)
                + "\n\n是否把这些新增条目合并进来？\n（只新增，不会覆盖你已有的数据）",
            ):
                for kind, (seed_only, new_keys) in extra.items():
                    try:
                        cur = md_store.load_rows(kind)
                        for k in new_keys:
                            cur.append(seed_only[k])
                        md_store.save_rows(kind, cur)
                        self._log(f"已合并 {md_store.TITLES[kind]} {len(new_keys)} 条", "ok")
                    except Exception as e:
                        self._log(f"合并 {md_store.TITLES[kind]} 失败：{e}", "error")

        self._refresh_md_count()

    def _refresh_md_count(self):
        try:
            products, corps, errors = md_store.load_indexes()
        except Exception:
            return
        if errors:
            for e in errors:
                self._log(f"主数据：{e}", "error")
            return
        self._log(f"主数据：产品 {len(products)} 条、法人 {len(corps)} 条", "info")

    def _open_masterdata(self, kind):
        if self.worker and self.worker.is_alive():
            messagebox.showinfo("请稍候", "正在转换中，转换完成后再修改主数据。")
            return
        win = self._md_win.get(kind)
        if win is not None and win.winfo_exists():
            win.deiconify()
            win.lift()
            win.focus_force()
            return
        if masterdata.MasterDataWindow is None:
            messagebox.showerror("不可用", "主数据界面加载失败（缺少 tkinter）。")
            return
        win = masterdata.MasterDataWindow(self.root, kind, on_change=self._on_md_changed)
        self._md_win[kind] = win

    def _on_md_changed(self):
        """主数据被改动后立即刷新一次统计（不重复弹窗）。"""
        try:
            products, corps, _ = md_store.load_indexes()
            self._log(f"主数据已更新：产品 {len(products)} 条、法人 {len(corps)} 条", "info")
        except Exception:
            pass

    def _open_md_folder(self):
        d = md_store.data_dir()
        try:
            os.makedirs(d, exist_ok=True)
        except Exception as e:
            messagebox.showerror("打不开", str(e))
            return
        self._open_path(d)

    def _about(self):
        messagebox.showinfo(
            "关于",
            "UDI XML 转 Excel\n\n"
            "支持：销售出库单（终端接口）、生产入库单（关联关系）\n"
            "输入：xml / zip / 文件夹\n\n"
            f"主数据目录：\n{md_store.data_dir()}",
        )

    def _help(self):
        messagebox.showinfo(
            "使用说明",
            "1. 把 xml、zip 或文件夹拖进窗口（或用按钮添加）\n"
            "2. 点「开始转换」\n\n"
            "销售出库单：每个 DN号 一个 Excel，含「销售出库明细」「校验」两个 sheet。\n"
            "成品名称、规格、供货者、客户名称来自主数据，\n"
            "未登记的编码会留空并在「校验」里列出来。\n\n"
            "主数据怎么维护：\n"
            "菜单「主数据」→ 产品/法人 → 上传 Excel 批量导入，\n"
            "或直接双击表格里的单元格修改。",
        )

    # ------------------------------------------------------------ 布局

    def _build_ui(self):
        self.root.title("UDI XML 转 Excel")
        self.root.geometry("860x700")
        self.root.minsize(760, 600)

        style = ttk.Style()
        try:
            style.theme_use("vista")
        except tk.TclError:
            pass

        self._build_menu()

        # 顶部标题
        head = ttk.Frame(self.root, padding=(16, 14, 16, 6))
        head.pack(fill="x")
        ttk.Label(head, text="UDI XML 转 Excel", font=("Microsoft YaHei UI", 15, "bold")).pack(anchor="w")
        ttk.Label(
            head,
            text="支持销售出库单、生产入库单的 xml / zip / 文件夹。"
                 "销售出库单会按 DI码 补成品名称与规格、按法人编码补供货者与客户名称。",
            foreground="#666",
        ).pack(anchor="w", pady=(3, 0))

        # 拖拽区
        drop_wrap = ttk.Frame(self.root, padding=(16, 8, 16, 0))
        drop_wrap.pack(fill="both", expand=False)
        self.drop = tk.Frame(drop_wrap, bg="#F5F8FF", highlightthickness=2,
                             highlightbackground="#B8CCF0", height=120)
        self.drop.pack(fill="x")
        self.drop.pack_propagate(False)
        self.drop_label = tk.Label(
            self.drop,
            text=("把 xml 文件、zip 压缩包或文件夹拖到这里" if DND_OK
                  else "点击下方按钮添加文件（当前环境不支持拖拽）"),
            bg="#F5F8FF", fg="#3B5BA5", font=("Microsoft YaHei UI", 11),
        )
        self.drop_label.pack(expand=True)

        if DND_OK:
            for w in (self.drop, self.drop_label):
                w.drop_target_register(DND_FILES)
                w.dnd_bind("<<Drop>>", self._on_drop)
            self.drop_label.configure(
                text="把 xml 文件、zip 压缩包或文件夹拖到这里\n也可以点下面的按钮选择",
                justify="center",
            )

        # 按钮行
        btn_row = ttk.Frame(self.root, padding=(16, 10, 16, 0))
        btn_row.pack(fill="x")
        ttk.Button(btn_row, text="添加文件…", command=self._add_files).pack(side="left")
        ttk.Button(btn_row, text="添加文件夹…", command=self._add_folder).pack(side="left", padx=(8, 0))
        ttk.Button(btn_row, text="清空列表", command=self._clear).pack(side="left", padx=(8, 0))
        self.count_label = ttk.Label(btn_row, text="已添加 0 个文件", foreground="#666")
        self.count_label.pack(side="right")

        # 文件列表
        list_wrap = ttk.Frame(self.root, padding=(16, 10, 16, 0))
        list_wrap.pack(fill="both", expand=True)
        cols = ("file", "type", "status")
        self.tree = ttk.Treeview(list_wrap, columns=cols, show="headings", height=9)
        self.tree.heading("file", text="文件")
        self.tree.heading("type", text="识别类型")
        self.tree.heading("status", text="状态")
        self.tree.column("file", width=430, anchor="w")
        self.tree.column("type", width=190, anchor="w")
        self.tree.column("status", width=200, anchor="w")
        vsb = ttk.Scrollbar(list_wrap, orient="vertical", command=self.tree.yview)
        self.tree.configure(yscrollcommand=vsb.set)
        self.tree.pack(side="left", fill="both", expand=True)
        vsb.pack(side="right", fill="y")
        self.tree.tag_configure("ok", foreground="#1B7F3B")
        self.tree.tag_configure("err", foreground="#C0392B")
        self.tree.tag_configure("warn", foreground="#B9770E")

        # 输出目录
        out_row = ttk.Frame(self.root, padding=(16, 10, 16, 0))
        out_row.pack(fill="x")
        ttk.Label(out_row, text="输出到：").pack(side="left")
        self.out_var = tk.StringVar(value=os.path.join(APP_DIR, "输出结果"))
        ttk.Entry(out_row, textvariable=self.out_var).pack(side="left", fill="x", expand=True, padx=(6, 6))
        ttk.Button(out_row, text="更改…", command=self._pick_out).pack(side="left")
        ttk.Button(out_row, text="打开", command=self._open_out).pack(side="left", padx=(6, 0))

        # 操作行
        act = ttk.Frame(self.root, padding=(16, 12, 16, 0))
        act.pack(fill="x")
        self.btn_start = ttk.Button(act, text="开始转换", command=self._start)
        self.btn_start.pack(side="left")
        self.btn_cancel = ttk.Button(act, text="取消", command=self._cancel, state="disabled")
        self.btn_cancel.pack(side="left", padx=(8, 0))
        self.progress = ttk.Progressbar(act, mode="determinate", maximum=100)
        self.progress.pack(side="left", fill="x", expand=True, padx=(12, 0))
        self.pct_label = ttk.Label(act, text="", width=6, anchor="e")
        self.pct_label.pack(side="left", padx=(6, 0))

        # 日志
        log_wrap = ttk.Frame(self.root, padding=(16, 10, 16, 14))
        log_wrap.pack(fill="both", expand=False)
        self.log = tk.Text(log_wrap, height=7, wrap="word", bg="#FAFAFA",
                           relief="solid", borderwidth=1, font=("Consolas", 9))
        lsb = ttk.Scrollbar(log_wrap, orient="vertical", command=self.log.yview)
        self.log.configure(yscrollcommand=lsb.set)
        self.log.pack(side="left", fill="both", expand=True)
        lsb.pack(side="right", fill="y")
        for tag, color in (("info", "#333"), ("ok", "#1B7F3B"),
                           ("warn", "#B9770E"), ("error", "#C0392B")):
            self.log.tag_configure(tag, foreground=color)
        self.log.configure(state="disabled")
        self._log(f"就绪。{'拖拽' if DND_OK else '按钮'}添加文件后点「开始转换」。", "info")

    # ------------------------------------------------------------ 日志 / 进度

    def _log(self, text, level="info"):
        self.log.configure(state="normal")
        self.log.insert("end", text + "\n", level)
        self.log.see("end")
        self.log.configure(state="disabled")

    def _poll_queue(self):
        """主线程定时取出工作线程的消息，避免跨线程操作界面。"""
        try:
            while True:
                kind, payload = self.queue.get_nowait()
                if kind == "log":
                    self._log(*payload)
                elif kind == "progress":
                    done, total, name = payload
                    pct = int(done * 100 / total) if total else 0
                    self.progress["value"] = pct
                    self.pct_label.configure(text=f"{pct}%")
                elif kind == "row":
                    idx, text, tag = payload
                    iid = f"row{idx}"
                    if self.tree.exists(iid):
                        self.tree.set(iid, "status", text)
                        self.tree.item(iid, tags=(tag,))
                elif kind == "done":
                    self._on_done(payload)
                elif kind == "fatal":
                    messagebox.showerror("出错了", payload)
        except queue.Empty:
            pass
        self.root.after(80, self._poll_queue)

    # ------------------------------------------------------------ 添加文件

    def _on_drop(self, event):
        try:
            paths = self.root.tk.splitlist(event.data)
        except Exception:
            paths = [event.data]
        self._add_paths(list(paths))

    def _add_files(self):
        paths = filedialog.askopenfilenames(
            title="选择 xml 或 zip 文件",
            filetypes=[("XML 和压缩包", "*.xml *.XML *.zip *.ZIP"), ("所有文件", "*.*")],
        )
        if paths:
            self._add_paths(list(paths))

    def _add_folder(self):
        d = filedialog.askdirectory(title="选择包含 xml / zip 的文件夹")
        if d:
            self._add_paths([d])

    def _add_paths(self, paths):
        if self.worker and self.worker.is_alive():
            messagebox.showinfo("请稍候", "正在转换中，请等待完成或点「取消」。")
            return
        try:
            new_sources, errors = src_mod.expand(paths)
        except Exception as e:
            messagebox.showerror("添加失败", f"读取文件时出错：\n{e}")
            return

        existing = {(s.display_name, id(s)) for s in self.sources}
        added = 0
        for s in new_sources:
            if (s.display_name, id(s)) in existing:
                continue
            self.sources.append(s)
            added += 1

        for e in errors:
            self._log("！ " + e, "warn")

        self._refresh_tree()
        if added:
            self._log(f"已添加 {added} 个文件（压缩包已展开）。", "info")
        elif not errors:
            self._log("没有新增可处理的 xml。", "warn")
        if errors and not added:
            messagebox.showwarning("没有可处理的文件", "\n".join(errors[:8]))

    def _refresh_tree(self):
        self.tree.delete(*self.tree.get_children())
        for i, s in enumerate(self.sources):
            self.tree.insert("", "end", iid=f"row{i}",
                             values=(s.display_name, "待识别", "待处理"))
        self.count_label.configure(text=f"已添加 {len(self.sources)} 个文件")

    def _clear(self):
        if self.worker and self.worker.is_alive():
            messagebox.showinfo("请稍候", "正在转换中，无法清空。")
            return
        self.sources.clear()
        self.results = []
        self._refresh_tree()
        self.progress["value"] = 0
        self.pct_label.configure(text="")

    # ------------------------------------------------------------ 输出目录

    def _pick_out(self):
        d = filedialog.askdirectory(title="选择输出目录")
        if d:
            self.out_var.set(d)

    def _open_out(self):
        d = self.out_var.get().strip()
        if not d:
            return
        if not os.path.isdir(d):
            messagebox.showinfo("目录不存在", "还没有输出结果，或目录已被移动。")
            return
        self._open_path(d)

    def _open_path(self, d):
        try:
            if os.name == "nt":
                os.startfile(d)  # noqa: S606
            elif sys.platform == "darwin":
                subprocess.Popen(["open", d])
            else:
                subprocess.Popen(["xdg-open", d])
        except Exception as e:
            messagebox.showerror("打不开", str(e))

    # ------------------------------------------------------------ 转换

    def _start(self):
        if self.worker and self.worker.is_alive():
            return
        if not self.sources:
            messagebox.showinfo("没有文件", "请先添加要转换的 xml 或压缩包。")
            return
        out_dir = self.out_var.get().strip()
        if not out_dir:
            messagebox.showinfo("缺少输出目录", "请先选择输出目录。")
            return

        self.cancel_flag.clear()
        self.btn_start.configure(state="disabled")
        self.btn_cancel.configure(state="normal")
        self.progress["value"] = 0
        self.pct_label.configure(text="0%")
        self.log.configure(state="normal")
        self.log.delete("1.0", "end")
        self.log.configure(state="disabled")
        self._log(f"输出目录：{out_dir}", "info")

        for i in range(len(self.sources)):
            iid = f"row{i}"
            if self.tree.exists(iid):
                self.tree.set(iid, "status", "处理中…")
                self.tree.item(iid, tags=())

        sources = list(self.sources)
        self.worker = threading.Thread(
            target=self._run_worker, args=(sources, out_dir), daemon=True
        )
        self.worker.start()

    def _run_worker(self, sources, out_dir):
        try:
            results, _ = convert(
                sources, out_dir,
                progress=lambda d, t, n: self.queue.put(("progress", (d, t, n))),
                cancel=lambda: self.cancel_flag.is_set(),
                log=lambda m, lv="info": self.queue.put(("log", (m, "ok" if lv == "ok" else lv))),
            )
            # 逐行回填状态
            name_to_idx = {s.display_name: i for i, s in enumerate(sources)}
            for r in results:
                idx = name_to_idx.get(r.display_name)
                if idx is None:
                    continue
                if r.ok:
                    txt = f"完成 {r.info} {('·校验通过' if r.verify_ok else '·自检异常')}"
                    tag = "ok" if r.verify_ok else "warn"
                else:
                    txt = "失败：" + r.error[:60]
                    tag = "err"
                self.queue.put(("row", (idx, txt, tag)))
            self.queue.put(("done", results))
        except Exception:
            self.queue.put(("fatal", "转换过程中出现意外错误：\n\n" + traceback.format_exc()))

    def _cancel(self):
        self.cancel_flag.set()
        self.btn_cancel.configure(state="disabled")
        self._log("正在取消…", "warn")

    def _on_done(self, results):
        self.results = results
        self.btn_start.configure(state="normal")
        self.btn_cancel.configure(state="disabled")

        ok = [r for r in results if r.ok]
        bad = [r for r in results if not r.ok]
        warn_verify = [r for r in ok if not r.verify_ok]

        self._log("─" * 52, "info")
        self._log(f"完成：成功 {len(ok)} 个，失败 {len(bad)} 个。", "ok" if not bad else "warn")
        if warn_verify:
            self._log(f"注意：{len(warn_verify)} 个文件写盘自检未通过，请抽查。", "error")
        if bad:
            for r in bad[:5]:
                self._log(f"  失败：{r.display_name} — {r.error}", "error")

        if bad:
            messagebox.showwarning(
                "部分文件未成功",
                f"成功 {len(ok)} 个，失败 {len(bad)} 个。\n详见下方日志。",
            )
        elif ok:
            if messagebox.askyesno("转换完成", f"成功转换 {len(ok)} 个文件。\n现在打开输出目录吗？"):
                self._open_out()


# ---------------------------------------------------------------- 入口

def run(argv=None):
    argv = sys.argv if argv is None else argv

    # 自检模式：设置 UDI_SELFTEST=<输出目录> 时跑一次转换后退出（用于验证打包产物）
    selftest = os.environ.get("UDI_SELFTEST")
    if selftest:
        return _selftest(selftest)

    try:
        if DND_OK:
            root = TkinterDnD.Tk()
        else:
            root = tk.Tk()
    except Exception as e:
        # 极端情况：没有图形环境
        print(f"无法启动图形界面：{e}")
        return 1
    App(root)
    root.mainloop()
    return 0


def _selftest(out_dir):
    """不打开窗口，直接跑一遍完整转换并逐项断言。结果同时写到日志文件。

    exe 是 console=False，标准输出看不到，所以必须落盘：
        <out_dir>\\_selftest.log
    """
    import traceback

    from config import columns as C_BILL

    # 控制台可能是 GBK，统一成能容错的 utf-8 输出
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

    lines = []

    def say(msg=""):
        lines.append(str(msg))
        try:
            print(msg, flush=True)
        except Exception:
            pass

    def flush():
        try:
            os.makedirs(out_dir, exist_ok=True)
            with open(os.path.join(out_dir, "_selftest.log"), "w", encoding="utf-8") as f:
                f.write("\n".join(lines) + "\n")
        except Exception:
            pass

    say("=== 自检模式 ===")
    say("Python: " + sys.version.split()[0])
    say("冻结包: " + str(getattr(sys, "frozen", False)))
    say("拖拽支持: " + str(DND_OK))
    say("主数据目录: " + md_store.data_dir())
    try:
        import tkinter
        say(f"tkinter: {tkinter.TkVersion}")
    except Exception as e:
        say(f"tkinter 失败: {e}")

    data = os.environ.get("UDI_SELFTEST_IN", "")
    paths = [p for p in data.split(";") if p]
    if not paths:
        say("未提供测试输入（环境变量 UDI_SELFTEST_IN）")
        flush()
        return 1

    failures = []

    def check(cond, msg):
        say(("  [OK]  " if cond else "  [FAIL] ") + msg)
        if not cond:
            failures.append(msg)

    try:
        sources, errors = src_mod.expand(paths)
    except Exception:
        say("展开输入时异常：")
        say(traceback.format_exc())
        flush()
        return 1

    if not sources:
        say("没有展开出任何 xml，自检无意义。请检查 UDI_SELFTEST_IN 里的路径是否正确。")
        for e in errors:
            say("  ! " + e)
        flush()
        return 1

    try:
        say(f"展开出 {len(sources)} 个 xml，提示 {len(errors)} 条")
        for e in errors:
            say("  ! " + e)
        results, _ = convert(sources, out_dir, log=lambda m, lv="info": say("  " + m))
    except Exception:
        say("自检异常：")
        say(traceback.format_exc())
        flush()
        return 1

    say("--- 转换结果 ---")
    by_name = {}
    for r in results:
        by_name[r.display_name] = r
        if r.ok:
            say(f"  OK  {r.display_name} -> {os.path.basename(r.out_path)} | {r.info}")
        else:
            say(f"  ERR {r.display_name}: {r.error}")
            failures.append(f"{r.display_name} 转换失败：{r.error}")

    # ---- 断言（只在输入里出现对应样例时才检查）
    def find(sub):
        for name, r in by_name.items():
            if sub in name and r.ok:
                return r
        return None

    r = find("销售出库单_样例")
    if r:
        import openpyxl
        wb = openpyxl.load_workbook(r.out_path)
        ws = wb["销售出库明细"]
        rows = list(ws.iter_rows(values_only=True))
        header = list(rows[0]) if rows else []
        expect_head = [c[1] for c in C_BILL.BILL_COLUMNS]
        check(os.path.basename(r.out_path) == "DN20240001.xlsx",
              f"文件名用 DN号：{os.path.basename(r.out_path)}")
        check(wb.sheetnames == ["销售出库明细", "校验"], f"sheet 只有 2 个：{wb.sheetnames}")
        check(header == expect_head, f"表头 {header}")
        if len(rows) >= 3:
            first, second = rows[1], rows[2]
            check(all(first[i] for i in (3, 4, 6, 7)), "第 1 行主数据四项全部命中")
            check(second[2] == "06974783049999" and not second[3] and not second[7],
                  "第 2 行查不到主数据 → 留空")
            check(isinstance(first[5], str) and len(first[5]) == 30, "UDI码 仍是 30 位文本")
        vs = list(wb["校验"].iter_rows(min_row=2, values_only=True))
        rules = [row[1] for row in vs]
        check("M_PRODUCT_NOT_FOUND" in rules, "记了 M_PRODUCT_NOT_FOUND")
        check("M_CORP_NOT_FOUND" in rules, "记了 M_CORP_NOT_FOUND")
        wb.close()

    r = find("无DN号")
    if r:
        check(os.path.basename(r.out_path) == "销售出库单_无DN号.xlsx",
              f"DN号为空时回退源文件名：{os.path.basename(r.out_path)}")

    r = find("生产入库单_样例")
    if r:
        import openpyxl
        wb = openpyxl.load_workbook(r.out_path)
        check(wb.sheetnames == ["盒明细", "箱汇总", "码清单", "校验"],
              f"入库单 4 个 sheet 齐全：{wb.sheetnames}")
        check(wb["盒明细"].max_row - 1 == 5, f"盒明细 5 行（实际 {wb['盒明细'].max_row - 1}）")
        check(wb["箱汇总"].max_row - 1 == 2, f"箱汇总 2 行（实际 {wb['箱汇总'].max_row - 1}）")
        wb.close()

    try:
        p_index, c_index, md_err = md_store.load_indexes()
        check(not md_err, f"主数据可读：产品 {len(p_index)} 条、法人 {len(c_index)} 条")
        check(len(p_index) > 0, f"产品主数据非空（{len(p_index)} 条）")
        check(len(c_index) > 0, f"法人主数据非空（{len(c_index)} 条）")
    except Exception as e:
        check(False, f"主数据读取失败：{e}")

    if not [r for r in results if r.ok]:
        failures.append("没有任何文件转换成功")

    say("=== 自检" + ("通过" if not failures else f"失败（{len(failures)} 项）") + " ===")
    flush()
    return 0 if not failures else 1


if __name__ == "__main__":
    sys.exit(run())
