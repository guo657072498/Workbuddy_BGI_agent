"""Excel 写出：主通道 openpyxl，兜底通道零依赖 OOXML。

关键约定：类型为 "text" 的列，单元格必须是字符串且 number_format 为 "@"，
否则 30 位数字码会被 Excel 转成科学计数法、末尾精度丢失。
"""

import os
import sys
import zipfile
from datetime import date, datetime
from xml.sax.saxutils import escape

TEXT, INT, DATE, AUTO = "text", "int", "date", "auto"


# ---------------------------------------------------------------- 公共

def _coerce(value, kind):
    """把值整理成适合写入的类型。返回 (值, 是否文本)。"""
    if value is None:
        return "", kind == TEXT
    if kind == TEXT:
        # 强制字符串，且不走数字分支
        return ("" if value == "" else str(value)), True
    if isinstance(value, bool):
        return int(value), False
    if isinstance(value, (int, float)):
        return value, False
    if isinstance(value, (date, datetime)):
        return value, False
    if kind == INT:
        try:
            return int(float(value)), False
        except (TypeError, ValueError):
            return ("" if value == "" else str(value)), True
    if value == "":
        return "", True
    # auto：纯数字转数字，其余保留文本
    s = str(value)
    try:
        if s.isdigit() or (s.lstrip("-").isdigit()):
            return int(s), False
        f = float(s)
        return f, False
    except ValueError:
        return s, True


def _rows_for(columns, records):
    """把记录列表转换为二维值表（含表头）。"""
    header = [c[1] for c in columns]
    body = []
    for rec in records:
        body.append([_coerce(rec.get(c[0], ""), c[2]) for c in columns])
    return header, body


# ---------------------------------------------------------------- openpyxl 通道

def write_openpyxl(path, sheets, highlight=True):
    """sheets: [(sheet_name, columns, records), ...]"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)

    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="4472C4")
    err_fill = PatternFill("solid", fgColor="FFC7CE")
    center = Alignment(horizontal="center", vertical="center")

    for name, columns, records, *rest in sheets:
        ws = wb.create_sheet(name)
        header, body = _rows_for(columns, records)

        ws.append(header)
        for cell in ws[1]:
            cell.font = head_font
            cell.fill = head_fill
            cell.alignment = center

        for row in body:
            # 空串写成真正的空单元格，避免 Excel 里显示成 None 或空字符串
            ws.append([(None if v == "" else v) for v, _ in row])
            r = ws.max_row
            for ci, (val, is_text) in enumerate(row, 1):
                if is_text:
                    ws.cell(row=r, column=ci).number_format = "@"

        # 列宽 + 格式
        for ci, col in enumerate(columns, 1):
            width = col[3] if len(col) > 3 and col[3] else 14
            ws.column_dimensions[get_column_letter(ci)].width = width
            if col[2] == TEXT:
                for r in range(2, ws.max_row + 1):
                    ws.cell(row=r, column=ci).number_format = "@"
            elif col[2] == DATE:
                for r in range(2, ws.max_row + 1):
                    ws.cell(row=r, column=ci).number_format = "yyyy-mm-dd"

        ws.freeze_panes = "A2"
        if ws.max_row > 1:
            ws.auto_filter.ref = ws.dimensions

        # 校验 sheet：按级别标色
        if highlight and name == "校验" and records:
            for r in range(2, ws.max_row + 1):
                lvl = ws.cell(row=r, column=1).value
                if lvl == "错误":
                    for ci in range(1, len(columns) + 1):
                        ws.cell(row=r, column=ci).fill = err_fill

    wb.save(path)
    return path


# ---------------------------------------------------------------- 零依赖通道

_CT = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
<Default Extension="xml" ContentType="application/xml"/>
<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>
<Override PartName="/xl/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/>
{overrides}
</Types>"""

_RELS = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/>
</Relationships>"""

_WB = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"
 xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
<sheets>{sheets}</sheets></workbook>"""

_WB_RELS = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
{rels}</Relationships>"""

_STYLES = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
<fonts count="2"><font><sz val="11"/><name val="Calibri"/></font>
<font><b/><sz val="11"/><color rgb="FFFFFFFF"/><name val="Calibri"/></font></fonts>
<fills count="3"><fill><patternFill patternType="none"/></fill>
<fill><patternFill patternType="gray125"/></fill>
<fill><patternFill patternType="solid"><fgColor rgb="FF4472C4"/><bgColor indexed="64"/></patternFill></fill></fills>
<borders count="1"><border><left/><right/><top/><bottom/><diagonal/></border></borders>
<cellStyleXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/></cellStyleXfs>
<cellXfs count="3">
<xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/>
<xf numFmtId="0" fontId="1" fillId="2" borderId="0" xfId="0" applyFont="1" applyFill="1"><alignment horizontal="center"/></xf>
<xf numFmtId="49" fontId="0" fillId="0" borderId="0" xfId="0" applyNumberFormat="1"/>
</cellXfs>
<cellStyles count="1"><cellStyle name="Normal" xfId="0" builtinId="0"/></cellStyles>
</styleSheet>"""


def _col_letter(i):
    """1 -> A, 27 -> AA"""
    s = ""
    while i > 0:
        i, r = divmod(i - 1, 26)
        s = chr(65 + r) + s
    return s


def _cell_xml(ref, value, is_text):
    if value == "" or value is None:
        return ""
    if is_text:
        return f'<c r="{ref}" s="2" t="inlineStr"><is><t xml:space="preserve">{escape(str(value))}</t></is></c>'
    if isinstance(value, bool):
        return f'<c r="{ref}" t="b"><v>{int(value)}</v></c>'
    if isinstance(value, (int, float)):
        return f'<c r="{ref}"><v>{value}</v></c>'
    return f'<c r="{ref}" t="inlineStr"><is><t xml:space="preserve">{escape(str(value))}</t></is></c>'


def write_stdlib(path, sheets, highlight=True):
    """零依赖写出：纯标准库手写 xlsx（inlineStr，无 sharedStrings）。"""
    n = len(sheets)

    overrides = "\n".join(
        f'<Override PartName="/xl/worksheets/sheet{i}.xml" '
        f'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
        for i in range(1, n + 1)
    )
    sheet_entries = "".join(
        f'<sheet name="{escape(str(s[0]))}" sheetId="{i}" r:id="rId{i}"/>'
        for i, s in enumerate(sheets, 1)
    )
    wb_rels = "".join(
        f'<Relationship Id="rId{i}" '
        f'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" '
        f'Target="worksheets/sheet{i}.xml"/>'
        for i in range(1, n + 1)
    ) + (f'<Relationship Id="rId{n + 1}" '
         f'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" '
         f'Target="styles.xml"/>')

    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", _CT.format(overrides=overrides))
        z.writestr("_rels/.rels", _RELS)
        z.writestr("xl/workbook.xml", _WB.format(sheets=sheet_entries))
        z.writestr("xl/_rels/workbook.xml.rels", _WB_RELS.format(rels=wb_rels))
        z.writestr("xl/styles.xml", _STYLES)

        for idx, (name, columns, records, *rest) in enumerate(sheets, 1):
            header, body = _rows_for(columns, records)
            rows_xml = []
            # 表头
            cells = "".join(
                f'<c r="{_col_letter(ci)}1" s="1" t="inlineStr"><is><t>{escape(str(h))}</t></is></c>'
                for ci, h in enumerate(header, 1)
            )
            rows_xml.append(f'<row r="1">{cells}</row>')

            for ri, row in enumerate(body, 2):
                cells = "".join(
                    _cell_xml(f"{_col_letter(ci)}{ri}", v, is_text)
                    for ci, (v, is_text) in enumerate(row, 1)
                )
                rows_xml.append(f'<row r="{ri}">{cells}</row>')

            cols_xml = "".join(
                f'<col min="{i}" max="{i}" width="{columns[i - 1][3] if len(columns[i - 1]) > 3 and columns[i - 1][3] else 14}" customWidth="1"/>'
                for i in range(1, len(columns) + 1)
            )
            sheet_xml = (
                '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
                f'<sheetViews><sheetView workbookViewId="0"><pane ySplit="1" topLeftCell="A2" activePane="bottomLeft" state="frozen"/></sheetView></sheetViews>'
                f'<cols>{cols_xml}</cols>'
                f'<sheetData>{"".join(rows_xml)}</sheetData>'
                '</worksheet>'
            )
            z.writestr(f"xl/worksheets/sheet{idx}.xml", sheet_xml)

    return path


# ---------------------------------------------------------------- 统一入口

def write_workbook(path, sheets, engine="auto", highlight=True):
    """写出 xlsx。engine: auto / openpyxl / stdlib。返回实际使用的引擎。"""
    if engine in ("auto", "openpyxl"):
        try:
            write_openpyxl(path, sheets, highlight=highlight)
            return "openpyxl"
        except ImportError:
            if engine == "openpyxl":
                raise
            # 降级
    write_stdlib(path, sheets, highlight=highlight)
    return "stdlib"
